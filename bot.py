import sqlite3
import asyncio
from urllib.parse import urlparse
import ipaddress
from difflib import get_close_matches
from contextlib import contextmanager
from dataclasses import dataclass, field
import tempfile
import aiohttp
import discord
from discord import app_commands
from discord.ext import commands, tasks
import signal
import csv
import json
import io
import os
import logging
from dotenv import load_dotenv

# ====================================
# LOGGING CONFIGURATION (Standard Library - zero RAM overhead)
# ====================================
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s [%(levelname)s] %(name)s: %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger("DiscordBot")

# ====================================
# CONFIGURATION (frozen dataclass — fail-fast on start, zero runtime overhead)
# Fill your IDs and credentials in the .env file, NOT here!
# ====================================
load_dotenv()

@dataclass(frozen=True)
class Config:
    """Central bot configuration. Values loaded from .env, fallbacks provided."""
    token: str = field(default_factory=lambda: os.getenv("DISCORD_TOKEN", ""))
    webhook_url: str = field(default_factory=lambda: os.getenv("WEBHOOK_URL", ""))
    # Put your Discord Guild (Server) ID and Channel IDs in your .env file
    guild_id: int = field(default_factory=lambda: int(os.getenv("GUILD_ID", "0") or "0"))
    admin_channel_id: int = field(default_factory=lambda: int(os.getenv("ADMIN_CHANNEL_ID", "0") or "0"))
    status_channel_id: int = field(default_factory=lambda: int(os.getenv("STATUS_CHANNEL_ID", "0") or "0"))
    # Can accept ANY SQLite database file path via DB_PATH in .env
    db_path: str = field(default_factory=lambda: os.getenv("DB_PATH", "database.db"))
    link_check_concurrency: int = 3   # Max concurrent link checks
    link_check_timeout: int = 8       # Timeout per request (seconds)
    per_page: int = 10                # Pagination items per page
    max_nazwa: int = 100              # Max product name length
    max_url: int = 2048               # Max URL length

config = Config()

# ====================================
# OPTIONAL DEPENDENCIES
# ====================================
try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# ====================================
# BOT INSTANCE
# ====================================
intents = discord.Intents.default()
intents.message_content = True

class CustomBot(commands.Bot):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.bot_uruchomiony_flag = False
        self._zamykanie_w_toku = False

    async def setup_hook(self):
        # Register POSIX signals (Linux VPS support)
        loop = asyncio.get_running_loop()
        for sig in (signal.SIGINT, signal.SIGTERM):
            try:
                loop.add_signal_handler(sig, lambda: asyncio.create_task(self.obsluga_wylaczenia()))
            except NotImplementedError:
                # Windows does not support add_signal_handler in asyncio loop
                pass

    async def czysc_kanal_statusu(self):
        """Purge status channel to remove old status messages."""
        kanal = self.get_channel(config.status_channel_id)
        if kanal:
            try:
                await kanal.purge(limit=100)
            except Exception as e:
                logger.warning(f"Error purging status channel: {e}")

    async def wyslij_status_webhook(self, tytul: str, opis: str, kolor: int):
        """Clean channel and send offline status via Webhook."""
        if not config.webhook_url or 'webhooks' not in config.webhook_url:
            return

        await self.czysc_kanal_statusu()

        payload = {
            "embeds": [{
                "title": tytul,
                "description": opis,
                "color": kolor
            }]
        }
        try:
            timeout = aiohttp.ClientTimeout(total=5)
            async with aiohttp.ClientSession(timeout=timeout) as session:
                await session.post(config.webhook_url, json=payload)
        except Exception as e:
            logger.warning(f"Error sending webhook: {e}")

    async def obsluga_wylaczenia(self):
        """Handles graceful shutdown and offline status notification."""
        if self._zamykanie_w_toku:
            return
        self._zamykanie_w_toku = True

        logger.info("🔴 Shutdown procedure captured! Sending offline status...")
        await self.wyslij_status_webhook(
            tytul="🔴 Bot has been shut down",
            opis="The bot has switched to **offline** mode and stopped responding.",
            kolor=15158332  # Red
        )
        await self.close()

bot = CustomBot(command_prefix="!", intents=intents)

# ====================================
# DATABASE — Context Manager (guarantees close + auto-commit/rollback)
# ====================================
@contextmanager
def get_db():
    conn = sqlite3.connect(config.db_path)
    conn.execute('PRAGMA journal_mode = WAL')
    conn.execute('PRAGMA synchronous = NORMAL')
    conn.execute('PRAGMA cache_size = -2000')   # Cache limit ~2MB RAM
    conn.execute('PRAGMA temp_store = MEMORY')   # Temp tables in RAM
    conn.execute('PRAGMA busy_timeout = 5000')   # Wait 5s on lock instead of crashing
    try:
        yield conn
    except Exception:
        conn.rollback()
        raise
    else:
        conn.commit()
    finally:
        conn.close()

# ====================================
# SCHEMA MIGRATIONS (versioned, idempotent, safe upgrade)
# ====================================
MIGRATIONS = [
    # v1 (idx 0): base table
    """CREATE TABLE IF NOT EXISTS produkty (
        nazwa TEXT PRIMARY KEY,
        link TEXT NOT NULL
    )""",
    # v2 (idx 1): stats column (compatibility with older databases)
    "ALTER TABLE produkty ADD COLUMN wywolania INTEGER DEFAULT 0",
    # v3 (idx 2): index on LOWER(nazwa) — accelerates case-insensitive search
    "CREATE INDEX IF NOT EXISTS idx_nazwa_lower ON produkty(LOWER(nazwa))",
    # v4 (idx 3): index on calls — accelerates /top
    "CREATE INDEX IF NOT EXISTS idx_wywolania ON produkty(wywolania DESC)",
    # v5 (idx 4): audit log — tracks who changed what and when
    """CREATE TABLE IF NOT EXISTS audit_log (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        timestamp TEXT DEFAULT (datetime('now', 'localtime')),
        user_id INTEGER NOT NULL,
        user_name TEXT NOT NULL,
        action TEXT NOT NULL,
        target_name TEXT,
        details TEXT
    )""",
]

def init_db():
    with get_db() as conn:
        conn.execute("CREATE TABLE IF NOT EXISTS schema_version (version INTEGER DEFAULT 0)")
        row = conn.execute("SELECT version FROM schema_version").fetchone()
        if not row:
            conn.execute("INSERT INTO schema_version (version) VALUES (0)")
            current_ver = 0
        else:
            current_ver = row[0]

        for i in range(current_ver, len(MIGRATIONS)):
            try:
                conn.execute(MIGRATIONS[i])
                logger.info(f"Migration {i + 1}/{len(MIGRATIONS)} OK")
            except sqlite3.OperationalError as e:
                err = str(e).lower()
                if "duplicate column" in err or "already exists" in err:
                    logger.debug(f"Migration {i + 1} skipped (already exists)")
                else:
                    raise

        conn.execute("UPDATE schema_version SET version = ?", (len(MIGRATIONS),))

# ====================================
# VALIDATION & SECURITY
# ====================================
def _escape_like(text: str) -> str:
    """Escapes LIKE special characters (%, _, \\) to prevent LIKE injection in autocomplete."""
    return text.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")

def jest_poprawnym_url(url: str) -> bool:
    """URL validation with SSRF protection — blocks private/reserved IPs using stdlib ipaddress."""
    try:
        result = urlparse(url)
        if result.scheme not in ('http', 'https') or not result.netloc:
            return False
        host = result.netloc.lower().split(':')[0]
        if not host or '.' not in host:
            return False
        try:
            addr = ipaddress.ip_address(host)
            if addr.is_private or addr.is_loopback or addr.is_link_local or addr.is_reserved:
                return False
        except ValueError:
            if host in ('localhost', 'localtest.me', 'lvh.me', 'vcap.me'):
                return False
        return True
    except Exception:
        return False

def waliduj_nazwe(nazwa: str) -> str | None:
    """Returns error message or None if product name is valid."""
    n = nazwa.strip()
    if len(n) < 1:
        return "Product name cannot be empty."
    if len(n) > config.max_nazwa:
        return f"Name cannot exceed {config.max_nazwa} characters (provided {len(n)})."
    return None

def waliduj_url(url: str) -> str | None:
    """Returns error message or None if URL is valid."""
    if len(url) > config.max_url:
        return f"URL cannot exceed {config.max_url} characters."
    if not jest_poprawnym_url(url):
        return "Invalid URL — required format `http://` or `https://`."
    return None

# ====================================
# DATABASE FUNCTIONS
# ====================================
def _db_podpowiadaj(current: str):
    with get_db() as conn:
        escaped = _escape_like(current.lower())
        cursor = conn.execute(
            "SELECT nazwa FROM produkty WHERE LOWER(nazwa) LIKE ? ESCAPE '\\' LIMIT 25",
            (f"%{escaped}%",)
        )
        return [row[0] for row in cursor.fetchall()]

def _db_szukaj(produkt: str):
    """Searches for a product. Returns (result_tuple, suggestions_list)."""
    with get_db() as conn:
        cursor = conn.execute(
            'SELECT nazwa, link, wywolania FROM produkty WHERE LOWER(nazwa) = ? LIMIT 1',
            (produkt.lower(),)
        )
        wynik = cursor.fetchone()
        if wynik:
            conn.execute(
                'UPDATE produkty SET wywolania = wywolania + 1 WHERE LOWER(nazwa) = ?',
                (produkt.lower(),)
            )
            nazwa, link, wywolania = wynik
            return (nazwa, link, wywolania + 1), []

        cursor = conn.execute('SELECT nazwa FROM produkty')
        wszystkie = [row[0] for row in cursor.fetchall()]
        sugestie = get_close_matches(produkt, wszystkie, n=3, cutoff=0.5)
        return None, sugestie

def _db_dodaj(nazwa: str, url: str):
    with get_db() as conn:
        try:
            conn.execute('INSERT INTO produkty (nazwa, link, wywolania) VALUES (?, ?, 0)', (nazwa, url))
            return "dodano"
        except sqlite3.IntegrityError:
            conn.execute('UPDATE produkty SET link = ? WHERE LOWER(nazwa) = ?', (url, nazwa.lower()))
            return "zaktualizowano"

def _db_usun(nazwa: str):
    with get_db() as conn:
        row = conn.execute(
            'SELECT nazwa FROM produkty WHERE LOWER(nazwa) = ?', (nazwa.lower(),)
        ).fetchone()
        if row:
            conn.execute('DELETE FROM produkty WHERE LOWER(nazwa) = ?', (nazwa.lower(),))
            return row[0]
        return None

def _db_edytuj(stara_nazwa: str, nowa_nazwa: str | None, nowy_link: str | None):
    with get_db() as conn:
        row = conn.execute(
            'SELECT nazwa, link FROM produkty WHERE LOWER(nazwa) = ?', (stara_nazwa.lower(),)
        ).fetchone()
        if not row:
            return False, f"Product `{stara_nazwa}` does not exist in the database."

        realna_nazwa, stary_link = row
        docelowa_nazwa = nowa_nazwa.strip() if nowa_nazwa else realna_nazwa
        docelowy_link = nowy_link.strip() if nowy_link else stary_link

        if docelowa_nazwa.lower() != stara_nazwa.lower():
            if conn.execute(
                'SELECT 1 FROM produkty WHERE LOWER(nazwa) = ?', (docelowa_nazwa.lower(),)
            ).fetchone():
                return False, f"Product named `{docelowa_nazwa}` already exists."

        conn.execute(
            'UPDATE produkty SET nazwa = ?, link = ? WHERE LOWER(nazwa) = ?',
            (docelowa_nazwa, docelowy_link, stara_nazwa.lower())
        )
        return True, f"Updated: **{docelowa_nazwa}** -> <{docelowy_link}>"

def _db_pobierz_strone(page: int, per_page: int):
    with get_db() as conn:
        total = conn.execute('SELECT COUNT(*) FROM produkty').fetchone()[0]
        offset = page * per_page
        items = conn.execute(
            'SELECT nazwa, link, COALESCE(wywolania, 0) FROM produkty ORDER BY nazwa ASC LIMIT ? OFFSET ?',
            (per_page, offset)
        ).fetchall()
        return total, items

def _db_top(limit: int = 10):
    with get_db() as conn:
        return conn.execute(
            'SELECT nazwa, link, wywolania FROM produkty ORDER BY wywolania DESC, nazwa ASC LIMIT ?',
            (limit,)
        ).fetchall()

def _db_statystyki():
    with get_db() as conn:
        total_items, total_searches = conn.execute(
            'SELECT COUNT(*), COALESCE(SUM(wywolania), 0) FROM produkty'
        ).fetchone()
    db_size = os.path.getsize(config.db_path) if os.path.exists(config.db_path) else 0
    return total_items, total_searches, db_size

def _db_lista_dla_checkera():
    with get_db() as conn:
        return conn.execute('SELECT nazwa, link FROM produkty ORDER BY nazwa ASC').fetchall()

# ====================================
# AUDIT LOG
# ====================================
def _db_audit(user_id: int, user_name: str, action: str, target: str, details: str = ""):
    with get_db() as conn:
        conn.execute(
            'INSERT INTO audit_log (user_id, user_name, action, target_name, details) VALUES (?, ?, ?, ?, ?)',
            (user_id, user_name, action, target, details)
        )

def _db_pobierz_audit(limit: int = 15):
    with get_db() as conn:
        return conn.execute(
            'SELECT timestamp, user_name, action, target_name, details FROM audit_log ORDER BY id DESC LIMIT ?',
            (limit,)
        ).fetchall()

# ====================================
# AUTOCOMPLETE & COOLDOWN
# ====================================
async def autocomplete_products(interaction: discord.Interaction, current: str):
    wyniki = await asyncio.to_thread(_db_podpowiadaj, current)
    return [
        app_commands.Choice(name=nazwa, value=nazwa)
        for nazwa in wyniki
    ]

def dynamiczny_cooldown(interaction: discord.Interaction) -> app_commands.Cooldown | None:
    if interaction.user.guild_permissions.administrator:
        return None
    return app_commands.Cooldown(3, 10.0)

# ====================================
# UI VIEWS (Buttons & Pagination)
# ====================================
class ProduktView(discord.ui.View):
    def __init__(self, url: str):
        super().__init__(timeout=None)
        self.add_item(discord.ui.Button(
            label="Open link ↗",
            url=url,
            style=discord.ButtonStyle.link
        ))

class PotwierdzUsuniecie(discord.ui.View):
    def __init__(self, nazwa: str, user_id: int, user_name: str):
        super().__init__(timeout=30)
        self.nazwa = nazwa
        self.user_id = user_id
        self.user_name = user_name
        self.message: discord.Message | None = None

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.user_id:
            await interaction.response.send_message("❌ This button is not for you.", ephemeral=True)
            return False
        return True

    @discord.ui.button(label="🗑️ Confirm deletion", style=discord.ButtonStyle.danger)
    async def potwierdz(self, interaction: discord.Interaction, button: discord.ui.Button):
        realna = await asyncio.to_thread(_db_usun, self.nazwa)
        if realna:
            await asyncio.to_thread(_db_audit, self.user_id, self.user_name, "usun", realna)
            await interaction.response.edit_message(
                content=f"🗑️ Deleted from database: **{realna}**", view=None
            )
        else:
            await interaction.response.edit_message(
                content=f"❌ Not found in database: `{self.nazwa}`", view=None
            )

    @discord.ui.button(label="Cancel", style=discord.ButtonStyle.secondary)
    async def anuluj(self, interaction: discord.Interaction, button: discord.ui.Button):
        await interaction.response.edit_message(content="↩️ Deletion cancelled.", view=None)

    async def on_timeout(self):
        for item in self.children:
            item.disabled = True
        if self.message:
            try:
                await self.message.edit(
                    content=f"⏰ Time to confirm deletion of `{self.nazwa}` has expired.", view=self
                )
            except Exception:
                pass

class ListaPaginacja(discord.ui.View):
    def __init__(self, author_id: int, total_count: int, per_page: int = 10):
        super().__init__(timeout=120)
        self.author_id = author_id
        self.total_count = total_count
        self.per_page = per_page
        self.current_page = 0
        self.total_pages = max(1, (total_count - 1) // per_page + 1)
        self.message: discord.Message | None = None
        self.aktualizuj_stan_przyciskow()

    def stworz_embed(self, items: list) -> discord.Embed:
        embed = discord.Embed(
            title="📦 Product List",
            color=discord.Color.blue()
        )
        opis = "".join(
            f"• **{nazwa}**: <{url}> *(Searches: {wywolania})*\n"
            for nazwa, url, wywolania in items
        )
        embed.description = opis if opis else "*No products on this page.*"
        embed.set_footer(
            text=f"Page {self.current_page + 1} of {self.total_pages} | Total items: {self.total_count}"
        )
        return embed

    def aktualizuj_stan_przyciskow(self):
        self.prev_button.disabled = (self.current_page == 0)
        self.next_button.disabled = (self.current_page >= self.total_pages - 1)

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if interaction.user.id != self.author_id:
            await interaction.response.send_message("❌ This pagination belongs to another user.", ephemeral=True)
            return False
        return True

    async def on_timeout(self):
        for item in self.children:
            item.disabled = True
        if self.message:
            try:
                await self.message.edit(view=self)
            except Exception:
                pass

    @discord.ui.button(label="◀ Previous", style=discord.ButtonStyle.secondary)
    async def prev_button(self, interaction: discord.Interaction, button: discord.ui.Button):
        if self.current_page > 0:
            self.current_page -= 1
            total, items = await asyncio.to_thread(_db_pobierz_strone, self.current_page, self.per_page)
            self.total_count = total
            self.total_pages = max(1, (total - 1) // self.per_page + 1)
            self.aktualizuj_stan_przyciskow()
            await interaction.response.edit_message(embed=self.stworz_embed(items), view=self)

    @discord.ui.button(label="Next ▶", style=discord.ButtonStyle.secondary)
    async def next_button(self, interaction: discord.Interaction, button: discord.ui.Button):
        if self.current_page < self.total_pages - 1:
            self.current_page += 1
            total, items = await asyncio.to_thread(_db_pobierz_strone, self.current_page, self.per_page)
            self.total_count = total
            self.total_pages = max(1, (total - 1) // self.per_page + 1)
            self.aktualizuj_stan_przyciskow()
            await interaction.response.edit_message(embed=self.stworz_embed(items), view=self)

# ====================================
# SLASH COMMANDS — SEARCH & STATS
# ====================================
@bot.tree.command(name="search", description="Search for a product link")
@app_commands.checks.dynamic_cooldown(dynamiczny_cooldown)
@app_commands.autocomplete(product=autocomplete_products)
async def search(interaction: discord.Interaction, product: str):
    wynik, sugestie = await asyncio.to_thread(_db_szukaj, product)

    if wynik:
        nazwa, url, wywolania = wynik
        embed = discord.Embed(
            title=f"🔗 {nazwa}",
            description=f"Link: <{url}>",
            color=discord.Color.green()
        )
        embed.set_footer(text=f"Views: {wywolania}")
        embed.timestamp = discord.utils.utcnow()
        view = ProduktView(url)
        await interaction.response.send_message(embed=embed, view=view, ephemeral=True)
    elif sugestie:
        lista = "\n".join(f"• `{s}`" for s in sugestie)
        await interaction.response.send_message(
            f"❌ Not found: `{product}`\n\n💡 **Did you mean:**\n{lista}",
            ephemeral=True
        )
    else:
        await interaction.response.send_message(f"❌ No link found for: `{product}`", ephemeral=True)

@bot.tree.command(name="list", description="Display all products saved in database")
@app_commands.checks.dynamic_cooldown(dynamiczny_cooldown)
async def list_cmd(interaction: discord.Interaction):
    total_count, items = await asyncio.to_thread(_db_pobierz_strone, 0, config.per_page)

    if total_count == 0:
        await interaction.response.send_message("📁 Database is currently empty.", ephemeral=True)
        return

    if total_count <= config.per_page:
        embed = discord.Embed(
            title="📦 Product List",
            color=discord.Color.blue()
        )
        opis = "".join(f"• **{nazwa}**: <{url}> *(Searches: {wywolania})*\n" for nazwa, url, wywolania in items)
        embed.description = opis
        embed.set_footer(text=f"Total items: {total_count}")
        await interaction.response.send_message(embed=embed, ephemeral=True)
    else:
        view = ListaPaginacja(interaction.user.id, total_count, per_page=config.per_page)
        await interaction.response.send_message(embed=view.stworz_embed(items), view=view, ephemeral=True)
        view.message = await interaction.original_response()

@bot.tree.command(name="top", description="Display top 10 most searched products")
@app_commands.checks.dynamic_cooldown(dynamiczny_cooldown)
async def top(interaction: discord.Interaction):
    produkty = await asyncio.to_thread(_db_top, 10)
    if not produkty:
        await interaction.response.send_message("📁 Database is currently empty.", ephemeral=True)
        return

    embed = discord.Embed(
        title="🔥 Top 10 Most Searched Products",
        color=discord.Color.gold()
    )
    opis = "".join(
        f"**{idx}.** **{nazwa}** — `{wywolania}` searches\n   └ <{url}>\n"
        for idx, (nazwa, url, wywolania) in enumerate(produkty, start=1)
    )
    embed.description = opis
    embed.timestamp = discord.utils.utcnow()
    await interaction.response.send_message(embed=embed, ephemeral=True)

@bot.tree.command(name="stats", description="Database summary and bot statistics")
@app_commands.checks.dynamic_cooldown(dynamiczny_cooldown)
async def stats(interaction: discord.Interaction):
    total_items, total_searches, db_size = await asyncio.to_thread(_db_statystyki)
    kb_size = db_size / 1024

    embed = discord.Embed(
        title="📊 Database Statistics",
        color=discord.Color.blue()
    )
    embed.add_field(name="📦 Products count", value=f"`{total_items}`", inline=True)
    embed.add_field(name="🔍 Total searches", value=f"`{total_searches}`", inline=True)
    embed.add_field(name="💾 Database size", value=f"`{kb_size:.1f} KB`", inline=True)
    embed.timestamp = discord.utils.utcnow()
    await interaction.response.send_message(embed=embed, ephemeral=True)

# ====================================
# SLASH COMMANDS — ADMINISTRATION
# ====================================
@bot.tree.command(name="add", description="Add or update a product in database (Admin)")
@app_commands.checks.has_permissions(administrator=True)
async def add(interaction: discord.Interaction, name: str, url: str):
    err = waliduj_nazwe(name)
    if err:
        await interaction.response.send_message(f"⚠️ {err}", ephemeral=True)
        return

    err = waliduj_url(url)
    if err:
        await interaction.response.send_message(f"⚠️ {err}", ephemeral=True)
        return

    name_s = name.strip()
    url_s = url.strip()
    status = await asyncio.to_thread(_db_dodaj, name_s, url_s)

    await asyncio.to_thread(
        _db_audit, interaction.user.id, str(interaction.user), status, name_s, url_s
    )

    if status == "dodano":
        await interaction.response.send_message(f"✅ Added to database: **{name_s}** -> <{url_s}>", ephemeral=True)
    else:
        await interaction.response.send_message(f"🔄 Updated link: **{name_s}** -> <{url_s}>", ephemeral=True)

@bot.tree.command(name="delete", description="Delete a product from database (Admin)")
@app_commands.checks.has_permissions(administrator=True)
@app_commands.autocomplete(name=autocomplete_products)
async def delete_cmd(interaction: discord.Interaction, name: str):
    view = PotwierdzUsuniecie(name.strip(), interaction.user.id, str(interaction.user))
    await interaction.response.send_message(
        f"⚠️ Are you sure you want to delete **{name.strip()}** from the database?",
        view=view,
        ephemeral=True
    )
    view.message = await interaction.original_response()

@bot.tree.command(name="edit", description="Edit product name or link without losing statistics (Admin)")
@app_commands.checks.has_permissions(administrator=True)
@app_commands.autocomplete(product=autocomplete_products)
async def edit(
    interaction: discord.Interaction,
    product: str,
    new_url: str | None = None,
    new_name: str | None = None
):
    if not new_url and not new_name:
        await interaction.response.send_message("⚠️ Provide a new name or a new link to update.", ephemeral=True)
        return

    if new_name:
        err = waliduj_nazwe(new_name)
        if err:
            await interaction.response.send_message(f"⚠️ {err}", ephemeral=True)
            return

    if new_url:
        err = waliduj_url(new_url)
        if err:
            await interaction.response.send_message(f"⚠️ {err}", ephemeral=True)
            return

    sukces, msg = await asyncio.to_thread(_db_edytuj, product, new_name, new_url)

    if sukces:
        zmiany = []
        if new_name:
            zmiany.append(f"name: {product} → {new_name.strip()}")
        if new_url:
            zmiany.append(f"link: {new_url.strip()}")
        await asyncio.to_thread(
            _db_audit, interaction.user.id, str(interaction.user),
            "edit", product, "; ".join(zmiany)
        )
        await interaction.response.send_message(f"✅ {msg}", ephemeral=True)
    else:
        await interaction.response.send_message(f"❌ {msg}", ephemeral=True)

@bot.tree.command(name="audit", description="Show recent database changes (Admin)")
@app_commands.checks.has_permissions(administrator=True)
async def audit(interaction: discord.Interaction, limit: int = 15):
    wpisy = await asyncio.to_thread(_db_pobierz_audit, min(limit, 25))
    if not wpisy:
        await interaction.response.send_message("📋 No entries in audit log.", ephemeral=True)
        return

    embed = discord.Embed(
        title="📋 Recent Database Changes",
        color=discord.Color.greyple()
    )

    ikony = {"dodaj": "➕", "dodano": "➕", "zaktualizowano": "🔄", "usun": "🗑️", "edytuj": "✏️", "import": "📥"}
    opis = ""
    for ts, user, action, target, details in wpisy:
        ikona = ikony.get(action, "📝")
        linia = f"{ikona} `{ts}` **{user}** → {action}: `{target}`"
        if details:
            linia += f"\n  └ {details}"
        linia += "\n"
        if len(opis) + len(linia) > 4000:
            break
        opis += linia

    embed.description = opis
    embed.timestamp = discord.utils.utcnow()
    await interaction.response.send_message(embed=embed, ephemeral=True)

@bot.tree.command(name="shutdown", description="Safely shutdown the bot and send offline status (Admin)")
@app_commands.checks.has_permissions(administrator=True)
async def shutdown(interaction: discord.Interaction):
    await interaction.response.send_message("🔴 Shutting down bot...", ephemeral=True)
    logger.info("Bot shutdown requested via /shutdown command.")
    await bot.obsluga_wylaczenia()

# ====================================
# EXPORT & IMPORT
# ====================================
def _generuj_eksport(format_typu: str):
    with get_db() as conn:
        if format_typu == "json":
            produkty = conn.execute(
                'SELECT nazwa, link, COALESCE(wywolania, 0) FROM produkty ORDER BY nazwa ASC'
            ).fetchall()
        else:
            produkty = conn.execute(
                'SELECT nazwa, link FROM produkty ORDER BY nazwa ASC'
            ).fetchall()

    if format_typu == "xlsx":
        if not HAS_OPENPYXL:
            return None, len(produkty), "Missing openpyxl library"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"
        ws.append(["nazwa", "link"])

        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)

        for row_idx, (nazwa, url) in enumerate(produkty, start=2):
            ws.cell(row=row_idx, column=1, value=nazwa)
            cell_link = ws.cell(row=row_idx, column=2, value=url)
            cell_link.hyperlink = url
            cell_link.style = "Hyperlink"

        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 85

        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        return excel_buffer.getvalue(), len(produkty), None

    elif format_typu == "csv":
        csv_buffer = io.StringIO()
        writer = csv.writer(csv_buffer)
        writer.writerow(["nazwa", "link"])
        for nazwa, url in produkty:
            klikalny_link = f'=HYPERLINK("{url}"; "{url}")'
            writer.writerow([nazwa, klikalny_link])
        return csv_buffer.getvalue().encode('utf-8-sig'), len(produkty), None

    elif format_typu == "json":
        data = [{"nazwa": row[0], "link": row[1], "wywolania": row[2]} for row in produkty]
        json_bytes = json.dumps(data, indent=4, ensure_ascii=False).encode('utf-8')
        return json_bytes, len(produkty), None

    elif format_typu == "db":
        fd, temp_path = tempfile.mkstemp(suffix='.db')
        os.close(fd)
        try:
            bck = sqlite3.connect(temp_path)
            src = sqlite3.connect(config.db_path)
            src.backup(bck)
            bck.close()
            src.close()
            with open(temp_path, 'rb') as f:
                data = f.read()
            return data, len(produkty), None
        finally:
            if os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except Exception:
                    pass

    return None, 0, "Unsupported format"

@bot.tree.command(name="export", description="Export product database (Admin only)")
@app_commands.checks.has_permissions(administrator=True)
@app_commands.choices(format=[
    app_commands.Choice(name="xlsx (.xlsx - Formatted Excel)", value="xlsx"),
    app_commands.Choice(name="csv (.csv - Standard CSV)", value="csv"),
    app_commands.Choice(name="json (.json - Full backup with counters)", value="json"),
    app_commands.Choice(name="db (.db - SQLite Copy)", value="db")
])
async def export_cmd(interaction: discord.Interaction, format: str):
    await interaction.response.defer(ephemeral=True)

    file_bytes, count, err = await asyncio.to_thread(_generuj_eksport, format)
    if err:
        await interaction.followup.send(f"❌ {err}", ephemeral=True)
        return

    discord_file = discord.File(io.BytesIO(file_bytes), filename=f"products.{format}")

    await interaction.followup.send(
        content=f"📦 Generated database file in **.{format.upper()}** format (Items: `{count}`):",
        file=discord_file,
        ephemeral=True
    )

def _db_import_csv(tekst: str):
    reader = csv.reader(io.StringIO(tekst))
    wiersze = list(reader)

    if wiersze and wiersze[0] and wiersze[0][0].strip().lower() in ('nazwa', 'name'):
        wiersze = wiersze[1:]

    with get_db() as conn:
        dodane = 0
        zaktualizowane = 0
        pominiete = 0

        for wiersz in wiersze:
            if len(wiersz) >= 2:
                nazwa = wiersz[0].strip()
                link = wiersz[1].strip()

                if link.startswith('=HYPERLINK("'):
                    if '";' in link:
                        link = link.split('";')[0].replace('=HYPERLINK("', '')
                    elif '",' in link:
                        link = link.split('",')[0].replace('=HYPERLINK("', '')

                if not nazwa or len(nazwa) > config.max_nazwa or not jest_poprawnym_url(link):
                    pominiete += 1
                    continue

                if conn.execute(
                    'SELECT 1 FROM produkty WHERE LOWER(nazwa) = ? LIMIT 1', (nazwa.lower(),)
                ).fetchone():
                    conn.execute('UPDATE produkty SET link = ? WHERE LOWER(nazwa) = ?', (link, nazwa.lower()))
                    zaktualizowane += 1
                else:
                    conn.execute('INSERT INTO produkty (nazwa, link, wywolania) VALUES (?, ?, 0)', (nazwa, link))
                    dodane += 1

        return dodane, zaktualizowane, pominiete

@bot.tree.command(name="import", description="Import product database from CSV file (Admin only)")
@app_commands.checks.has_permissions(administrator=True)
async def import_data(interaction: discord.Interaction, file: discord.Attachment):
    if not file.filename.endswith('.csv'):
        await interaction.response.send_message("❌ Only **.csv** format is supported (columns: name, link).", ephemeral=True)
        return

    await interaction.response.defer(ephemeral=True)

    try:
        zawartosc = await file.read()
        tekst = zawartosc.decode('utf-8-sig')
        dodane, zaktualizowane, pominiete = await asyncio.to_thread(_db_import_csv, tekst)

        await asyncio.to_thread(
            _db_audit, interaction.user.id, str(interaction.user),
            "import", file.filename,
            f"added: {dodane}, updated: {zaktualizowane}, skipped: {pominiete}"
        )

        msg = f"✅ Import completed. Added **{dodane}** new, updated **{zaktualizowane}**."
        if pominiete > 0:
            msg += f"\n⚠️ Skipped **{pominiete}** items (invalid name or URL)."
        await interaction.followup.send(msg, ephemeral=True)
    except Exception as e:
        logger.error(f"Error during CSV import: {e}", exc_info=True)
        await interaction.followup.send(f"❌ Error parsing file: {e}", ephemeral=True)

# ====================================
# LINK HEALTH CHECKER
# ====================================
async def _sprawdz_link(session: aiohttp.ClientSession, nazwa: str, url: str, semaphore: asyncio.Semaphore):
    async with semaphore:
        for attempt in range(2):
            try:
                async with session.head(url, allow_redirects=True) as resp:
                    if resp.status == 405:
                        async with session.get(url, allow_redirects=True) as resp_get:
                            if resp_get.status >= 400:
                                return (nazwa, url, f"HTTP {resp_get.status}")
                    elif resp.status >= 500 and attempt == 0:
                        await asyncio.sleep(2)
                        continue
                    elif resp.status >= 400:
                        return (nazwa, url, f"HTTP {resp.status}")
                return None

            except asyncio.TimeoutError:
                if attempt == 0:
                    await asyncio.sleep(2)
                    continue
                return (nazwa, url, "Timeout")
            except aiohttp.ClientError:
                if attempt == 0:
                    await asyncio.sleep(2)
                    continue
                return (nazwa, url, "Connection error")
            except Exception:
                return (nazwa, url, "Unknown error")
    return None

async def wykonaj_sprawdzanie_linkow() -> int:
    produkty = await asyncio.to_thread(_db_lista_dla_checkera)
    if not produkty:
        return 0

    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
    }

    semaphore = asyncio.Semaphore(config.link_check_concurrency)
    timeout = aiohttp.ClientTimeout(total=config.link_check_timeout)
    connector = aiohttp.TCPConnector(ssl=False, limit=10)

    async with aiohttp.ClientSession(headers=headers, timeout=timeout, connector=connector) as session:
        zadania = [_sprawdz_link(session, nazwa, url, semaphore) for nazwa, url in produkty]
        wyniki = await asyncio.gather(*zadania)

    uszkodzone = [w for w in wyniki if w is not None]

    if uszkodzone:
        kanal_admin = bot.get_channel(config.admin_channel_id)
        if kanal_admin:
            naglowek = f"🚨 **[Link Checker Report] Found {len(uszkodzone)} broken links!**\n\n"
            linie = [f"• **{nazwa}**: <{url}>\n  └ *Reason:* `{powod}`" for nazwa, url, powod in uszkodzone]
            wiadomosc = naglowek
            for linia in linie:
                if len(wiadomosc) + len(linia) + 2 > 1900:
                    await kanal_admin.send(wiadomosc)
                    wiadomosc = ""
                wiadomosc += linia + "\n"
            if wiadomosc.strip():
                await kanal_admin.send(wiadomosc)

    return len(uszkodzone)

@tasks.loop(hours=24)
async def sprawdzaj_linki_w_tle():
    await wykonaj_sprawdzanie_linkow()

@sprawdzaj_linki_w_tle.before_loop
async def przed_sprawdzaniem():
    await bot.wait_until_ready()

@bot.tree.command(name="check_links", description="Run immediate link health check (Admin only)")
@app_commands.checks.has_permissions(administrator=True)
async def check_links(interaction: discord.Interaction):
    await interaction.response.defer(ephemeral=True)
    znaleziono = await wykonaj_sprawdzanie_linkow()
    if znaleziono == 0:
        await interaction.followup.send("✅ All links in database are working properly!", ephemeral=True)
    else:
        await interaction.followup.send(f"⚠️ Found **{znaleziono}** broken links. Report sent to admin channel.", ephemeral=True)

# ====================================
# ERROR HANDLING
# ====================================
@bot.tree.error
async def on_app_command_error(interaction: discord.Interaction, error: app_commands.AppCommandError):
    if isinstance(error, app_commands.CommandOnCooldown):
        msg = f"⏳ Slow down! Command ready in **{error.retry_after:.1f}s**."
    elif isinstance(error, app_commands.MissingPermissions):
        msg = "❌ Missing permissions to execute this command."
    else:
        logger.error(f"[Command Error] {error}", exc_info=True)
        msg = "❌ An unexpected error occurred while executing the command."

    if interaction.response.is_done():
        await interaction.followup.send(msg, ephemeral=True)
    else:
        await interaction.response.send_message(msg, ephemeral=True)

# ====================================
# EVENTS & SYNC
# ====================================
@bot.event
async def on_ready():
    await asyncio.to_thread(init_db)
    guild = discord.Object(id=config.guild_id)

    try:
        logger.info("Copying and syncing application commands...")
        bot.tree.copy_global_to(guild=guild)
        synced = await bot.tree.sync(guild=guild)
        logger.info(f"✅ SYNCED {len(synced)} COMMANDS TO GUILD: {[c.name for c in synced]}")
    except Exception as e:
        logger.error(f"❌ Command sync error: {e}")

    if not sprawdzaj_linki_w_tle.is_running():
        sprawdzaj_linki_w_tle.start()
        logger.info("🔍 Background link checking loop started.")

    logger.info(f"✅ Bot ready as {bot.user}")

    if not bot.bot_uruchomiony_flag:
        await bot.czysc_kanal_statusu()
        kanal_status = bot.get_channel(config.status_channel_id)
        if kanal_status:
            embed = discord.Embed(
                title="🟢 Bot successfully started",
                description="Bot is online and ready to accept commands.",
                color=discord.Color.green()
            )
            await kanal_status.send(embed=embed)
            bot.bot_uruchomiony_flag = True

if __name__ == "__main__":
    if not config.token:
        raise ValueError("❌ Missing DISCORD_TOKEN in .env file!")
    try:
        bot.run(config.token)
    except KeyboardInterrupt:
        logger.info("KeyboardInterrupt shutdown caught.")